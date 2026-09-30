"""Лист консультации в формате Word.

Документ, который можно распечатать, подписать и вложить в карту:
шапка, данные пациента, поля листа по шаблону специальности, подпись врача.
"""

from datetime import date, datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


def _ru(iso: str | None) -> str:
    if not iso:
        return ""
    try:
        return datetime.fromisoformat(iso).strftime("%d.%m.%Y")
    except ValueError:
        return iso


def _pair(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.add_run(f"{label}: ").bold = True
    p.add_run(value or "—")


def build(path: str, consultation: dict, template: dict, doctor: dict,
          organization: str = "") -> str:
    fields = consultation["fields"]
    patient = consultation.get("patient") or {}

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    for s in doc.sections:
        s.left_margin, s.right_margin = Cm(3), Cm(1.5)
        s.top_margin, s.bottom_margin = Cm(2), Cm(2)

    if organization:
        p = doc.add_paragraph(organization)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].bold = True

    h = doc.add_paragraph(template.get("sheet_title", "Лист консультации").upper())
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.runs[0].bold = True
    h.runs[0].font.size = Pt(14)

    meta = doc.add_paragraph(f"от {_ru(consultation.get('visit_date'))}")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _pair(doc, "Пациент", patient.get("name", ""))
    _pair(doc, "ИИН", patient.get("iin", ""))
    _pair(doc, "Дата рождения", _ru(patient.get("birth_date")))
    _pair(doc, "Врач", f"{doctor.get('name', '')}, {doctor.get('position', '')}".strip(", "))
    doc.add_paragraph()

    for spec in template["fields"]:
        f = fields.get(spec["id"])
        if not f:
            continue
        kind = f["kind"]

        if kind == "prescriptions":
            doc.add_paragraph().add_run(f"{f['title']}:").bold = True
            items = f.get("items") or []
            if not items:
                doc.add_paragraph("—")
                continue
            table = doc.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            for i, name in enumerate(["Препарат", "Доза", "Кратность", "Срок"]):
                cell = table.rows[0].cells[i]
                cell.text = name
                cell.paragraphs[0].runs[0].bold = True
            for item in items:
                row = table.add_row().cells
                row[0].text = item.get("drug", "")
                row[1].text = item.get("dose", "")
                row[2].text = item.get("frequency", "")
                row[3].text = item.get("duration", "")
            doc.add_paragraph()
            continue

        if kind == "diagnosis":
            icd = f.get("icd") or {}
            code = f" (МКБ-10: {icd['code']})" if icd.get("code") else ""
            _pair(doc, f["title"], (f.get("value") or "") + code if f.get("value") else "")
            continue

        if kind in ("date_future", "date_past"):
            value = _ru(f.get("date")) or f.get("value", "")
            _pair(doc, f["title"], value)
            preg = f.get("pregnancy")
            if preg:
                _pair(doc, "Срок беременности", f"{preg['weeks']} нед. {preg['days']} дн.")
                _pair(doc, "Предполагаемая дата родов", _ru(preg["edd"]))
            continue

        _pair(doc, f["title"], f.get("value", ""))

    doc.add_paragraph()
    sign = doc.add_paragraph()
    sign.add_run(f"Врач: {doctor.get('name', '')}")
    sign.add_run("\t\t\tПодпись: ________________")

    foot = doc.add_paragraph()
    run = foot.add_run(
        "\nЛист заполнен по записи приёма системой Хатшы, проверен и подтверждён врачом. "
        f"Сформирован {date.today():%d.%m.%Y}."
    )
    run.italic = True
    run.font.size = Pt(9)

    doc.save(path)
    return path
